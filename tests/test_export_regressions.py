import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from exporter import export_arp_xlsx, export_xlsx
from session import SessionManager


class ExportRoundTripTests(unittest.TestCase):
    def test_discovered_values_remain_text_in_every_session_sheet(self):
        session = SessionManager()
        session.add_port_scan({"name": "=1+1", "port": "#REF!", "vlan": 42})
        session.add_ip_snapshot({"hostname": "=device"},
                                [(15, "Domain", "=example", "standard")])
        session.add_mdns_scan([{"friendly": "=printer", "raw": "#N/A"}])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "session.xlsx"
            export_xlsx(session, path)
            workbook = load_workbook(path)
            try:
                for sheet, address, value in (
                    ("Port Scans", "C2", "=1+1"),
                    ("Port Scans", "D2", "#REF!"),
                    ("IP Snapshots", "C2", "=device"),
                    ("IP Snapshots", "D7", "=example"),
                    ("mDNS Discoveries", "C2", "=printer"),
                    ("mDNS Discoveries", "F2", "#N/A"),
                ):
                    with self.subTest(sheet=sheet, address=address):
                        cell = workbook[sheet][address]
                        self.assertEqual(cell.value, value)
                        self.assertEqual(cell.data_type, "s")
                self.assertEqual(workbook["Port Scans"]["G2"].value, 42)
                self.assertEqual(workbook["Port Scans"]["G2"].data_type, "n")
            finally:
                workbook.close()

    def test_control_characters_do_not_prevent_session_export(self):
        session = SessionManager()
        session.add_port_scan({"name": "switch\x00\x0bname\twith\nlines"})
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "session.xlsx"
            export_xlsx(session, path)
            workbook = load_workbook(path)
            try:
                self.assertEqual(workbook["Port Scans"]["C2"].value,
                                 "switchname\twith\nlines")
            finally:
                workbook.close()

    def test_arp_export_preserves_hostnames_as_text_and_ignores_missing_mac(self):
        results = [
            {"ip": "192.0.2.1", "mac": "00:11:22:33:44:55", "hostname": "=host\x00"},
            {"ip": "192.0.2.2", "mac": "", "hostname": "missing"},
        ]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "arp.xlsx"
            export_arp_xlsx(results, path)
            workbook = load_workbook(path)
            try:
                sheet = workbook["ARP"]
                self.assertEqual(sheet.max_row, 2)
                self.assertEqual(sheet["C2"].value, "=host")
                self.assertEqual(sheet["C2"].data_type, "s")
                self.assertEqual(sheet["A1"].value, "IP Address")
                self.assertEqual(sheet["B1"].value, "MAC Address")
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
